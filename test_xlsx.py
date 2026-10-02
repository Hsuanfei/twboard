# SPDX-License-Identifier: GPL-3.0-only
# Copyright (C) 2026 黃炫斐 (Mick Huang)
# 本檔案是「台股戰略產生器」的一部分：自由軟體，依 GNU GPL 第 3 版釋出，不附任何擔保，詳見 LICENSE。
"""0930a：這次分析的全部股票一次下載成 Excel（總覽＋每檔一張工作表），只用標準函式庫產生 .xlsx。"""
import io
import json
import os
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile

os.environ.setdefault("TWBOARD_CACHE_DIR", tempfile.mkdtemp(prefix="twboard-xl-"))
os.environ.setdefault("TWBOARD_DATA_DIR", tempfile.mkdtemp(prefix="twboard-xl-data-"))
import twboard as T           # noqa: E402
import twserve as S           # noqa: E402
import twxlsx as X            # noqa: E402

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def sheets_of(blob):
    """回傳 [(工作表名稱, [[儲存格文字或數字…]…])]，用標準函式庫讀，順便確認每個 XML 都能解析。"""
    z = zipfile.ZipFile(io.BytesIO(blob))
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            ET.fromstring(z.read(name))
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    names = [s.get("name") for s in wb.find("m:sheets", NS)]
    out = []
    for i, n in enumerate(names):
        ws = ET.fromstring(z.read("xl/worksheets/sheet%d.xml" % (i + 1)))
        rows = []
        for row in ws.find("m:sheetData", NS):
            cells = []
            for c in row:
                t = c.find("m:is/m:t", NS)
                v = c.find("m:v", NS)
                cells.append(t.text if t is not None else (float(v.text) if v is not None else None))
            rows.append(cells)
        out.append((n, rows))
    return out


class WriterTests(unittest.TestCase):
    def test_cells_and_names(self):
        blob = X.workbook([
            {"name": "0050 元大台灣50", "header": ["代號", "日期", "數值", "文字"],
             "rows": [["0050", "2026-09-29", 111.3, "a<b&c"], ["00403A", "2026-09-30", 7, None]], "notes": ["資料來源：測試"]},
            {"name": "a/b:c*d?[e]" + "x" * 40, "header": ["x"], "rows": []},
            {"name": "0050 元大台灣50", "header": ["y"], "rows": [[1]]},
        ], title="測試")
        got = sheets_of(blob)
        self.assertEqual([n for n, _ in got][0], "0050 元大台灣50")
        long = got[1][0]
        self.assertLessEqual(len(long), 31); self.assertFalse(set("[]:*?/\\") & set(long), long)
        self.assertEqual(got[2][0], "0050 元大台灣50(2)", "重複名稱自動加編號")
        rows = got[0][1]
        self.assertEqual(rows[1][0], "0050", "代號是文字，保留前導 0")
        self.assertEqual(rows[1][1], 46294.0, "ISO 日期存成 Excel 日期序號（2026-09-29）")
        self.assertEqual(rows[1][2:], [111.3, "a<b&c"])
        self.assertEqual(rows[2][:3], ["00403A", 46295.0, 7.0]); self.assertEqual(len(rows[2]), 3, "None 留空不寫儲存格")
        self.assertEqual(rows[-1], ["資料來源：測試"])
        with self.assertRaises(ValueError):
            X.workbook([])

    def test_openpyxl_can_read_if_installed(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("沒有安裝 openpyxl（程式本身不需要）")
        blob = X.workbook([{"name": "S", "header": ["日期", "收"], "rows": [["2026-09-29", 1.5]]}])
        ws = openpyxl.load_workbook(io.BytesIO(blob))["S"]
        self.assertEqual(ws.freeze_panes, "A2"); self.assertEqual(ws.auto_filter.ref, "A1:B2")
        self.assertEqual(ws["A2"].value.date().isoformat(), "2026-09-29"); self.assertEqual(ws["B2"].value, 1.5)


class PayloadTests(unittest.TestCase):
    def test_overview_and_one_sheet_per_stock(self):
        ps = [T.analyse(c, T.fetch_demo(c, days=160, seed=i, history_days=600), 30) for i, c in enumerate(["0050", "2330", "00403A"])]
        got = sheets_of(S.to_xlsx(ps))
        self.assertEqual([n for n, _ in got], ["總覽", "0050 測試樣本", "2330 測試樣本", "00403A 測試樣本"])
        ov = got[0][1]
        self.assertEqual(ov[0][:5], ["代號", "名稱", "市場", "股價最新日期", "收盤"])
        self.assertEqual([r[0] for r in ov[1:4]], ["0050", "2330", "00403A"])
        self.assertEqual(ov[1][4], ps[0]["quote"]["close"])
        self.assertIn(ov[1][ov[0].index("除權息確認")], ("已確認", "未取得"), "狀態用中文，不是程式代碼")
        header, rows = got[1][1][0], got[1][1][1:]
        self.assertEqual(header[:6], ["日期", "開盤", "最高", "最低", "收盤", "成交量(張)"])
        self.assertEqual(header[-1], "K 線型態")
        data = [r for r in rows if isinstance(r[0], float)]
        self.assertEqual(len(data), 30, "每檔的工作表是分析天數內的交易日")
        self.assertEqual(data[-1][4], ps[0]["series"]["close"][-1])
        self.assertTrue(any(isinstance(r[0], str) and r[0].startswith("資料來源") for r in rows), "表格下方附資料來源")


class ServerTests(unittest.TestCase):
    def setUp(self):
        S.DEMO_MODE = True
        self.server = S.ThreadingHTTPServer(("127.0.0.1", 0), S.Handler)
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.base = "http://127.0.0.1:%d" % self.server.server_address[1]

    def tearDown(self):
        self.server.shutdown(); self.server.server_close(); S.DEMO_MODE = False

    def analyse(self, code):
        req = urllib.request.Request(self.base + "/api/jobs", data=json.dumps({"code": code, "days": "30"}).encode(),
                                     headers={"Content-Type": "application/json", "Origin": self.base})
        job = json.loads(urllib.request.urlopen(req, timeout=30).read())["job_id"]
        for _ in range(300):
            j = json.loads(urllib.request.urlopen(self.base + "/api/jobs/" + job, timeout=30).read())["job"]
            if j["status"] in ("done", "error"):
                return j["data"]
            time.sleep(0.05)

    def test_download_all_selected(self):
        ids = [self.analyse(c)["snapshot_id"] for c in ("2330", "0056")]
        with urllib.request.urlopen(self.base + "/api/xlsx?snapshots=" + ",".join(ids), timeout=30) as r:
            blob = r.read()
            self.assertIn("spreadsheetml", r.headers["Content-Type"])
            cd = r.headers["Content-Disposition"]
        self.assertIn(urllib.parse.quote("個股資料_"), cd); self.assertIn(urllib.parse.quote("_2檔.xlsx"), cd)
        self.assertEqual([n for n, _ in sheets_of(blob)], ["總覽", "2330 測試樣本", "0056 測試樣本"])
        for bad in ("", "?snapshots=", "?snapshots=missing"):
            with self.assertRaises(urllib.error.HTTPError) as e:
                urllib.request.urlopen(self.base + "/api/xlsx" + bad, timeout=30)
            self.assertEqual(e.exception.code, 400)


if __name__ == "__main__":
    unittest.main()
